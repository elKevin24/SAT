import openpyxl
wb = openpyxl.load_workbook('docs/fuentes-datos/02_archivos_producto_entregables/Catalogo_Contenidos_Unicos_Portal_SAT.xlsx')
out = []
for s in wb.sheetnames:
    ws = wb[s]
    out.append(f'=== {s} === {ws.max_row}x{ws.max_column}')
    for r in range(1, min(12, ws.max_row+1)):
        row_vals = []
        for c in range(1, min(10, ws.max_column+1)):
            row_vals.append(ws.cell(r,c).value)
        out.append(f'{r}: {row_vals}')
    out.append('')
with open('inspect2_out.txt','w',encoding='utf-8') as f:
    f.write('\n'.join(out))
