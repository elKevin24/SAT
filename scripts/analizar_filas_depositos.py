import openpyxl

excel_path = 'c:/Users/busqu/Documents/GitHub/SAT/docs/Detalle de Contenido para Grupos de Interes.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb['Auxiliares Función Pública']

print(f"{'Fila':<5} | {'URL':<65} | {'Alm.Fisc':<8} | {'Almacenad':<9} | {'Dep.Aduan':<9} | {'Descripción SAT'}")
print("-" * 140)

for r in range(12, ws.max_row + 1):
    no = ws.cell(row=r, column=2).value
    url = str(ws.cell(row=r, column=3).value or '')
    alm_fisc = str(ws.cell(row=r, column=6).value or '').strip()
    almacenad = str(ws.cell(row=r, column=7).value or '').strip()
    dep_aduan = str(ws.cell(row=r, column=8).value or '').strip()
    desc = str(ws.cell(row=r, column=13).value or '').strip()
    
    # Si aplica a alguno de los 3
    if alm_fisc in ['SI', 'Si', 'si'] or almacenad in ['SI', 'Si', 'si'] or dep_aduan in ['SI', 'Si', 'si']:
        # Truncar url y desc para tabla
        url_short = url if len(url) < 65 else url[:62] + '...'
        desc_short = desc if len(desc) < 45 else desc[:42] + '...'
        print(f"{r:<5} | {url_short:<65} | {alm_fisc:<8} | {almacenad:<9} | {dep_aduan:<9} | {desc_short}")
