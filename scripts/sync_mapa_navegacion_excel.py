import openpyxl

file_path = 'docs/fuentes-datos/02_archivos_producto_entregables/Mapa_de_Navegacion_y_Descripciones_Portal_SAT.xlsx'
wb = openpyxl.load_workbook(file_path)

print("=== SINCRONIZANDO Mapa_de_Navegacion_y_Descripciones_Portal_SAT.xlsx ===")

# -------------------------------------------------------------
# 1. HOJA: Mapa de Navegación
# -------------------------------------------------------------
ws_mapa = wb['Mapa de Navegación']

# Corregir Fila 7 (NIT sin Obligaciones)
ws_mapa.cell(row=7, column=7, value=10) # Columna G = 10 trámites
ws_mapa.cell(row=7, column=4, value='Sin Obligaciones') # Nivel 2 canónico
ws_mapa.cell(row=7, column=5, value='—') # Nivel 3 canónico compactado
ws_mapa.cell(row=7, column=10, value='Contribuyentes > Sin Obligaciones > [Gestión]') # Miga

# Sincronizar fórmulas de porcentaje en todas las filas de datos
# Contribuyentes (Rows 7 to 17) -> Subtotal Row 18
for r in range(7, 18):
    ws_mapa.cell(row=r, column=8, value=f"=G{r}/$G$18") # % Nivel 1
    ws_mapa.cell(row=r, column=9, value=f"=G{r}/$G$38") # % Portal

# Subtotal Contribuyentes (Row 18)
ws_mapa.cell(row=18, column=7, value="=SUM(G7:G17)")
ws_mapa.cell(row=18, column=8, value="100.0%")
ws_mapa.cell(row=18, column=9, value="=G18/$G$38")

# Operadores de Comercio Exterior (Rows 19 to 24) -> Subtotal Row 25
for r in range(19, 25):
    ws_mapa.cell(row=r, column=8, value=f"=G{r}/$G$25")
    ws_mapa.cell(row=r, column=9, value=f"=G{r}/$G$38")

# Subtotal Comercio Exterior (Row 25)
ws_mapa.cell(row=25, column=7, value="=SUM(G19:G24)")
ws_mapa.cell(row=25, column=8, value="100.0%")
ws_mapa.cell(row=25, column=9, value="=G25/$G$38")

# Profesionales (Rows 26 to 30) -> Subtotal Row 31
for r in range(26, 31):
    ws_mapa.cell(row=r, column=8, value=f"=G{r}/$G$31")
    ws_mapa.cell(row=r, column=9, value=f"=G{r}/$G$38")

# Subtotal Profesionales (Row 31)
ws_mapa.cell(row=31, column=7, value="=SUM(G26:G30)")
ws_mapa.cell(row=31, column=8, value="100.0%")
ws_mapa.cell(row=31, column=9, value="=G31/$G$38")

# Entes Exentos (Rows 32 to 36) -> Subtotal Row 37
for r in range(32, 37):
    ws_mapa.cell(row=r, column=8, value=f"=G{r}/$G$37")
    ws_mapa.cell(row=r, column=9, value=f"=G{r}/$G$38")

# Subtotal Entes Exentos (Row 37)
ws_mapa.cell(row=37, column=7, value="=SUM(G32:G36)")
ws_mapa.cell(row=37, column=8, value="100.0%")
ws_mapa.cell(row=37, column=9, value="=G37/$G$38")

# Gran Total Portal (Row 38)
ws_mapa.cell(row=38, column=7, value="=G18+G25+G31+G37")
ws_mapa.cell(row=38, column=8, value="—")
ws_mapa.cell(row=38, column=9, value="100.0%")
ws_mapa.cell(row=38, column=10, value="715 Trámites 100% Mapeados")

print("Hoja 'Mapa de Navegación' sincronizada.")

# -------------------------------------------------------------
# 2. HOJA: Fichas por Rol y Categoría
# -------------------------------------------------------------
ws_fichas = wb['Fichas por Rol y Categoría']

# Corregir Fila 7 (NIT sin Obligaciones)
ws_fichas.cell(row=7, column=3, value='Sin Obligaciones')
ws_fichas.cell(row=7, column=4, value='—')
ws_fichas.cell(row=7, column=8, value=10) # Columna H = 10 trámites

# Fila 34 (Total)
ws_fichas.cell(row=34, column=2, value='TOTAL PORTAL WEB SAT')
ws_fichas.cell(row=34, column=3, value='4 Macro-Grupos (N1)')
ws_fichas.cell(row=34, column=4, value='27 Categorías Principales')
ws_fichas.cell(row=34, column=8, value="=SUM(H7:H33)") # Suma = 715

print("Hoja 'Fichas por Rol y Categoría' sincronizada.")

# -------------------------------------------------------------
# 3. HOJA: Resumen Arquitectura
# -------------------------------------------------------------
ws_resumen = wb['Resumen Arquitectura']

# Banners superiores (Row 6)
ws_resumen.cell(row=6, column=2, value="TOTAL CONTENIDOS\n715 Trámites")
ws_resumen.cell(row=6, column=3, value="CONTRIBUYENTES\n344 Trámites")
ws_resumen.cell(row=6, column=4, value="COMERCIO EXTERIOR\n245 Trámites")
ws_resumen.cell(row=6, column=6, value="PROFESIONALES\n47 Trámites")
ws_resumen.cell(row=6, column=7, value="ENTES EXENTOS\n79 Trámites")

# Tabla 1: Macro-Grupos N1
ws_resumen.cell(row=11, column=7, value=245)
ws_resumen.cell(row=12, column=7, value=344)
ws_resumen.cell(row=13, column=7, value=47)
ws_resumen.cell(row=14, column=7, value=79)
ws_resumen.cell(row=15, column=7, value="=SUM(G11:G14)")
ws_resumen.cell(row=15, column=8, value="100.0%")

# Tabla 2: Ciclo ATO
ws_resumen.cell(row=19, column=7, value=79)   # Empezar y registrarse
ws_resumen.cell(row=19, column=8, value="=G19/$G$24")
ws_resumen.cell(row=20, column=7, value=260)  # Operación y declaraciones
ws_resumen.cell(row=20, column=8, value="=G20/$G$24")
ws_resumen.cell(row=21, column=7, value=116)  # Consultas y herramientas
ws_resumen.cell(row=21, column=8, value="=G21/$G$24")
ws_resumen.cell(row=22, column=7, value=130)  # Modificaciones y cierre
ws_resumen.cell(row=22, column=8, value="=G22/$G$24")
ws_resumen.cell(row=23, column=7, value=130)  # Normativa y asistencia
ws_resumen.cell(row=23, column=8, value="=G23/$G$24")
ws_resumen.cell(row=24, column=7, value="=SUM(G19:G23)")
ws_resumen.cell(row=24, column=8, value="100.0%")

# Tabla 3: Tipos de Interacción
ws_resumen.cell(row=28, column=7, value=485)  # Guía Informativa / Texto
ws_resumen.cell(row=28, column=8, value="=G28/$G$32")
ws_resumen.cell(row=29, column=7, value=141)  # Trámite / Aplicativo en Línea
ws_resumen.cell(row=29, column=8, value="=G29/$G$32")
ws_resumen.cell(row=30, column=7, value=75)   # Consulta en Base de Datos
ws_resumen.cell(row=30, column=8, value="=G30/$G$32")
ws_resumen.cell(row=31, column=7, value=14)   # Descarga de Documento / Software
ws_resumen.cell(row=31, column=8, value="=G31/$G$32")
ws_resumen.cell(row=32, column=7, value="=SUM(G28:G31)")
ws_resumen.cell(row=32, column=8, value="100.0%")

print("Hoja 'Resumen Arquitectura' sincronizada.")

wb.save(file_path)
print(f"Libro {file_path} guardado y blindado al 100%.")
