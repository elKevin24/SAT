import json
import openpyxl

# 1. Inspeccionar en allTramites.json
with open('c:/Users/busqu/Documents/GitHub/SAT/src/data/allTramites.json', 'r', encoding='utf-8') as f:
    tramites = json.load(f)

depositos_en_json = [
    t for t in tramites
    if t.get('pillar') == 'comercio_exterior' and (
        'depósito' in t.get('categoria', '').lower() or
        'deposito' in t.get('categoria', '').lower() or
        'almacén' in t.get('categoria', '').lower() or
        'almacen' in t.get('categoria', '').lower() or
        'depósito' in t.get('subcategoria', '').lower() or
        'almacén' in t.get('subcategoria', '').lower() or
        'depósito' in t.get('tramite', '').lower() or
        'almacén' in t.get('tramite', '').lower()
    )
]

print(f"Total coincidencias en allTramites.json: {len(depositos_en_json)}")
for d in depositos_en_json[:10]:
    print(f"- ID: {d['id']} | Cat: {d['categoria']} | Subcat: {d['subcategoria']} | Trámite: {d['tramite']}")

# 2. Inspeccionar en docs/Detalle de Contenido para Grupos de Interes.xlsx
excel_path = 'c:/Users/busqu/Documents/GitHub/SAT/docs/Detalle de Contenido para Grupos de Interes.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
sheet_names = wb.sheetnames
print(f"\nHojas en {excel_path}: {sheet_names}")

ce_sheet = None
for s in sheet_names:
    if 'comercio' in s.lower() or 'exterior' in s.lower():
        ce_sheet = wb[s]
        break

if ce_sheet:
    print(f"\nFilas de Depósitos Aduaneros en hoja '{ce_sheet.title}':")
    header = [cell.value for cell in ce_sheet[1]]
    print("Cabeceras:", header[:7])
    
    count = 0
    for r in ce_sheet.iter_rows(min_row=2, values_only=True):
        row_str = " ".join([str(v) for v in r if v is not None]).lower()
        if 'depósito' in row_str or 'deposito' in row_str or 'almacen' in row_str or 'almacén' in row_str:
            count += 1
            if count <= 15:
                print(f"Row {count}: N2={r[0]} | N3={r[1]} | N4={r[2]} | Tema={r[3]} | Subtema={r[4]} | Nombre={r[5]}")
    print(f"Total filas encontradas en Excel: {count}")
