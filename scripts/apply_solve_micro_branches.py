import json
import openpyxl
import os

print("=== INICIANDO SANEAMIENTO DE MICRO-RAMAS PASARELA ===")

# 1. ACTUALIZAR allTramites.json
json_path = 'src/data/allTramites.json'
with open(json_path, 'r', encoding='utf-8') as f:
    tramites = json.load(f)

for t in tramites:
    cid = t.get('codigo_oficial') or t.get('id')
    
    # Caso 1: SAT-GES-0097 (Descarga de aplicaciones móviles SAT)
    if cid == 'SAT-GES-0097':
        print(f"Actualizando {cid}...")
        t['subcategoria'] = 'RTU Digital y Agencia Virtual'
        t['tema'] = '—'
        t['subtema'] = '—'
        t['nivel3_subcategoria'] = 'RTU Digital y Agencia Virtual'
        t['nivel4_tema'] = '—'
        t['nivel5_tramite'] = '—'
        t['orden_n3'] = 2
        t['orden_n4'] = '—'
        t['orden_n5'] = '—'
        t['migaBreadcrumb'] = 'Contribuyentes > Contribuyente General > RTU Digital y Agencia Virtual > Descarga de aplicaciones móviles oficiales de la SAT'

    # Caso 2: SAT-GES-0161 (Proyección educativa)
    elif cid == 'SAT-GES-0161':
        print(f"Actualizando {cid}...")
        t['categoria'] = 'Contribuyente General'
        t['subcategoria'] = 'Cultura tributaria y capacitación'
        t['tema'] = 'Programas de cultura tributaria'
        t['subtema'] = '—'
        t['nivel2_categoria'] = 'Contribuyente General'
        t['nivel3_subcategoria'] = 'Cultura tributaria y capacitación'
        t['nivel4_tema'] = 'Programas de cultura tributaria'
        t['nivel5_tramite'] = '—'
        t['orden_n3'] = 5
        t['orden_n4'] = 3
        t['orden_n5'] = '—'
        t['migaBreadcrumb'] = 'Contribuyentes > Contribuyente General > Cultura tributaria y capacitación > Programas de cultura tributaria > Proyección educativa'

    # Caso 3: SAT-GES-0318 (Traspaso de vehículos por notario - vista comprador/vendedor)
    elif cid == 'SAT-GES-0318':
        print(f"Actualizando {cid}...")
        t['categoria'] = 'Contribuyente General'
        t['subcategoria'] = 'Vehículos (como propietario)'
        t['tema'] = 'Traspaso de vehículos'
        t['subtema'] = '—'
        t['nivel2_categoria'] = 'Contribuyente General'
        t['nivel3_subcategoria'] = 'Vehículos (como propietario)'
        t['nivel4_tema'] = 'Traspaso de vehículos'
        t['nivel5_tramite'] = '—'
        t['orden_n3'] = 3
        t['orden_n4'] = 2
        t['orden_n5'] = '—'
        t['migaBreadcrumb'] = 'Contribuyentes > Contribuyente General > Vehículos (como propietario) > Traspaso de vehículos > Traspaso de vehículos por notario (vista del comprador/vendedor)'

with open(json_path, 'w', encoding='utf-8') as f:
    json.dump(tramites, f, ensure_ascii=False, indent=2)
print("allTramites.json actualizado con éxito.")

# 2. ACTUALIZAR database/sat_portal_nosql.json si existe
nosql_path = 'database/sat_portal_nosql.json'
if os.path.exists(nosql_path):
    with open(nosql_path, 'r', encoding='utf-8') as f:
        nosql = json.load(f)
    if 'tramites' in nosql:
        for t in nosql['tramites']:
            cid = t.get('codigo_oficial') or t.get('id')
            if cid == 'SAT-GES-0097':
                t['subcategoria'] = 'RTU Digital y Agencia Virtual'
                t['tema'] = '—'
                t['subtema'] = '—'
                t['orden_n4'] = '—'
                t['orden_n5'] = '—'
            elif cid == 'SAT-GES-0161':
                t['subcategoria'] = 'Cultura tributaria y capacitación'
                t['tema'] = 'Programas de cultura tributaria'
                t['subtema'] = '—'
                t['orden_n5'] = '—'
            elif cid == 'SAT-GES-0318':
                t['subcategoria'] = 'Vehículos (como propietario)'
                t['tema'] = 'Traspaso de vehículos'
                t['subtema'] = '—'
                t['orden_n5'] = '—'
        with open(nosql_path, 'w', encoding='utf-8') as f:
            json.dump(nosql, f, ensure_ascii=False, indent=2)
        print("sat_portal_nosql.json actualizado con éxito.")

# 3. ACTUALIZAR EXCEL MAESTRO
excel_path = 'docs/fuentes-datos/02_archivos_producto_entregables/Estructura_Final_Contenido_Portal_SAT_Actualizado.xlsx'
wb = openpyxl.load_workbook(excel_path)

for sheet_name in ['Matriz Maestra (715)', 'Contribuyentes (344)']:
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers = [cell.value for cell in ws[1]]
        col_id = headers.index('ID Gestión') + 1
        col_n2 = headers.index('Nivel 2') + 1
        col_o3 = headers.index('Orden N3') + 1
        col_n3 = headers.index('Nivel 3') + 1
        col_o4 = headers.index('Orden N4') + 1
        col_n4 = headers.index('Nivel 4') + 1
        col_o5 = headers.index('Orden N5') + 1
        col_n5 = headers.index('Nivel 5') + 1
        col_miga = headers.index('Ruta de Navegación (Miga de Pan)') + 1

        for row_idx in range(2, ws.max_row + 1):
            val_id = ws.cell(row=row_idx, column=col_id).value
            if val_id == 'SAT-GES-0097':
                ws.cell(row=row_idx, column=col_n3, value='RTU Digital y Agencia Virtual')
                ws.cell(row=row_idx, column=col_o4, value='—')
                ws.cell(row=row_idx, column=col_n4, value='—')
                ws.cell(row=row_idx, column=col_o5, value='—')
                ws.cell(row=row_idx, column=col_n5, value='—')
                ws.cell(row=row_idx, column=col_miga, value='Contribuyentes > Contribuyente General > RTU Digital y Agencia Virtual > Descarga de aplicaciones móviles oficiales de la SAT')
            elif val_id == 'SAT-GES-0161':
                ws.cell(row=row_idx, column=col_n2, value='Contribuyente General')
                ws.cell(row=row_idx, column=col_o3, value=5)
                ws.cell(row=row_idx, column=col_n3, value='Cultura tributaria y capacitación')
                ws.cell(row=row_idx, column=col_o4, value=3)
                ws.cell(row=row_idx, column=col_n4, value='Programas de cultura tributaria')
                ws.cell(row=row_idx, column=col_o5, value='—')
                ws.cell(row=row_idx, column=col_n5, value='—')
                ws.cell(row=row_idx, column=col_miga, value='Contribuyentes > Contribuyente General > Cultura tributaria y capacitación > Programas de cultura tributaria > Proyección educativa')
            elif val_id == 'SAT-GES-0318':
                ws.cell(row=row_idx, column=col_n2, value='Contribuyente General')
                ws.cell(row=row_idx, column=col_o3, value=3)
                ws.cell(row=row_idx, column=col_n3, value='Vehículos (como propietario)')
                ws.cell(row=row_idx, column=col_o4, value=2)
                ws.cell(row=row_idx, column=col_n4, value='Traspaso de vehículos')
                ws.cell(row=row_idx, column=col_o5, value='—')
                ws.cell(row=row_idx, column=col_n5, value='—')
                ws.cell(row=row_idx, column=col_miga, value='Contribuyentes > Contribuyente General > Vehículos (como propietario) > Traspaso de vehículos > Traspaso de vehículos por notario (vista del comprador/vendedor)')

wb.save(excel_path)
print(f"Libro {excel_path} actualizado y guardado.")

print("=== SANEAMIENTO COMPLETADO ===")
