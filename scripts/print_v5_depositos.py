import openpyxl

excel_path = 'c:/Users/busqu/Documents/GitHub/SAT/docs/Arbol_de_Navegacion_Portal_v5.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb['Operadores de comercio exterior']

# Buscar filas donde Nivel 3 sea Almacenes Fiscales, Almacenadoras o Depósitos Aduaneros
matches = []
for r in range(2, ws.max_row + 1):
    n2 = ws.cell(row=r, column=1).value
    n3 = ws.cell(row=r, column=2).value
    tema = ws.cell(row=r, column=3).value
    subtema = ws.cell(row=r, column=4).value
    tramite_orig = ws.cell(row=r, column=5).value
    tramite_sug = ws.cell(row=r, column=6).value
    desc = ws.cell(row=r, column=7).value
    url = ws.cell(row=r, column=8).value
    
    if n3 in ['Almacenes Fiscales', 'Almacenadoras', 'Depósitos Aduaneros']:
        matches.append((r, n2, n3, tema, subtema, tramite_orig, tramite_sug, desc, url))

print(f"Total registros en 'Operadores de comercio exterior' para estos 3 actores: {len(matches)}")
from collections import Counter
print("Por Actor (N3):", Counter(m[2] for m in matches))

for idx, m in enumerate(matches, start=1):
    print(f"\n[{idx}] Fila {m[0]} | N3: {m[2]} | Tema: {m[3]} | Subtema: {m[4]}")
    print(f"    Trámite Sugerido: {m[6]}")
    print(f"    Trámite Original: {m[5]}")
    print(f"    Descripción: {m[7]}")
    print(f"    URL: {m[8]}")
