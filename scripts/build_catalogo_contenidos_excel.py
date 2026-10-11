import json
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ==============================================================================
# 1. CARGA DE DATOS
# ==============================================================================
JSON_PATH = 'src/data/catalogoContenidosUnicos.json'
with open(JSON_PATH, 'r', encoding='utf-8') as f:
    catalogo = json.load(f)

TOTAL_ROWS = len(catalogo) # 683
LAST_ROW = TOTAL_ROWS + 1  # 684

print(f"Cargados {TOTAL_ROWS} contenidos únicos desde {JSON_PATH}")

# ==============================================================================
# 2. CONFIGURACIÓN INSTITUCIONAL SAT
# ==============================================================================
wb = openpyxl.Workbook()
wb.remove(wb.active) # Eliminar hoja por defecto

# Paleta de Colores SAT
BLUE_HEADER = '19324B'      # Azul Oscuro Título
NAVY_HEADER = '14649B'      # Azul SAT Primario
CYAN_ACCENT = '0284C7'      # Celeste Operaciones
GREEN_ACCENT = '059669'     # Verde Consultas
ORANGE_ACCENT = 'C25E00'    # Naranja Alerta / Brechas
PURPLE_ACCENT = '7C3AED'    # Morado Normativa
ZEBRA_FILL = 'F8FAFC'       # Gris tenue alternado
HIGHLIGHT_FILL = 'EFF6FF'   # Azul tenue transversal
BORDER_COLOR = 'CBD5E1'     # Borde neutro

# Fuentes Segoe UI
font_title = Font(name='Segoe UI', size=16, bold=True, color='FFFFFF')
font_subtitle = Font(name='Segoe UI', size=10, italic=True, color='FFFFFF')
font_header = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
font_body = Font(name='Segoe UI', size=9)
font_bold = Font(name='Segoe UI', size=9, bold=True)
font_link = Font(name='Segoe UI', size=9, color='0284C7', underline='single')

thin_border = Border(
    left=Side(style='thin', color=BORDER_COLOR),
    right=Side(style='thin', color=BORDER_COLOR),
    top=Side(style='thin', color=BORDER_COLOR),
    bottom=Side(style='thin', color=BORDER_COLOR)
)

header_border = Border(
    left=Side(style='thin', color='FFFFFF'),
    right=Side(style='thin', color='FFFFFF'),
    top=Side(style='medium', color='FFFFFF'),
    bottom=Side(style='medium', color='FFFFFF')
)

# ==============================================================================
# HOJA 1: RESUMEN EJECUTIVO Y DASHBOARD
# ==============================================================================
ws_resumen = wb.create_sheet(title="Resumen Ejecutivo")
ws_resumen.views.sheetView[0].showGridLines = True

# Banner Institucional
ws_resumen.merge_cells('B2:H3')
banner = ws_resumen['B2']
banner.value = "SUPERINTENDENCIA DE ADMINISTRACIÓN TRIBUTARIA — SAT GUATEMALA"
banner.font = font_title
banner.fill = PatternFill(start_color=BLUE_HEADER, end_color=BLUE_HEADER, fill_type='solid')
banner.alignment = Alignment(horizontal='center', vertical='center')

ws_resumen.merge_cells('B4:H4')
sub_banner = ws_resumen['B4']
sub_banner.value = "CATÁLOGO DESACOPLADO DE CONTENIDOS Y GESTIONES ÚNICAS (CONTENT HUB)"
sub_banner.font = font_subtitle
sub_banner.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type='solid')
sub_banner.alignment = Alignment(horizontal='center', vertical='center')

SHEET_DATA_NAME = f"Catálogo Único ({TOTAL_ROWS})"

# KPIs Dinámicos con fórmulas
kpis = [
    ("TOTAL CONTENIDOS ÚNICOS", f"=COUNTA('{SHEET_DATA_NAME}'!A2:A{LAST_ROW})", "Entidades Canónicas", "B6", "B7", NAVY_HEADER),
    ("APARICIONES EN EL ÁRBOL", f"=SUM('{SHEET_DATA_NAME}'!O2:O{LAST_ROW})", "Nodos de Navegación", "C6", "D7", CYAN_ACCENT),
    ("CONTENIDOS EXCLUSIVOS", f'=COUNTIF(\'{SHEET_DATA_NAME}\'!N2:N{LAST_ROW}, "NO")', "Aparecen 1 sola vez", "E6", "E7", GREEN_ACCENT),
    ("TRANSVERSALES POLIJERÁRQUICOS", f'=COUNTIF(\'{SHEET_DATA_NAME}\'!N2:N{LAST_ROW}, "SÍ")', "Múltiples audiencias", "F6", "G7", PURPLE_ACCENT),
    ("BRECHAS NORMATIVAS", f'=COUNTIF(\'{SHEET_DATA_NAME}\'!Q2:Q{LAST_ROW}, "SÍ")', "Propuestas Mesa Técnica", "H6", "H7", ORANGE_ACCENT),
]

for label, formula, sub, top_left, bot_right, color in kpis:
    ws_resumen.merge_cells(f"{top_left}:{bot_right}")
    cell = ws_resumen[top_left]
    cell.value = formula
    cell.font = Font(name='Segoe UI', size=13, bold=True, color='FFFFFF')
    cell.fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

# Tabla 1: Distribución por Etapa Ciclo de Vida
ws_resumen['B9'] = "1. DISTRIBUCIÓN POR ETAPA DEL CICLO DE VIDA DE LA GESTIÓN"
ws_resumen['B9'].font = Font(name='Segoe UI', size=11, bold=True, color='1E293B')

t1_headers = ["No.", "Etapa del Ciclo de Vida", "Descripción Operativa", "Total Contenidos", "% del Catálogo"]
for col_idx, h in enumerate(t1_headers, start=2):
    c = ws_resumen.cell(row=10, column=col_idx, value=h)
    c.font = font_header
    c.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type='solid')
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = header_border
ws_resumen.row_dimensions[10].height = 24

etapas_ato = [
    (1, "Operar y Declarar", "Rutina tributaria, declaraciones mensuales/anuales, pagos y transmisiones DUCA.", "Operar y Declarar"),
    (2, "Modificar y Cerrar", "Actualizaciones de RTU, traspasos, ceses prescritos y renovación anual de fianzas.", "Modificar y Cerrar"),
    (3, "Normativa y Asistencia", "Marco legal, resoluciones, criterios tributarios/aduaneros y capacitaciones virtuales.", "Normativa y Asistencia"),
    (4, "Consultas y Herramientas", "Verificadores abiertos, solvencia fiscal, consulta de NIT y selectivo aduanero.", "Consultas y Herramientas"),
    (5, "Empezar y Registrarse", "Inscripción inicial de NIT, padrones, carnés de auxiliares y habilitación de usuario.", "Empezar y Registrarse"),
]

for idx, (num, etapa, desc, label_exact) in enumerate(etapas_ato, start=11):
    ws_resumen.cell(row=idx, column=2, value=num).alignment = Alignment(horizontal='center')
    ws_resumen.cell(row=idx, column=3, value=etapa).font = font_bold
    ws_resumen.cell(row=idx, column=4, value=desc)
    c_tot = ws_resumen.cell(row=idx, column=5, value=f'=COUNTIF(\'{SHEET_DATA_NAME}\'!G2:G{LAST_ROW}, "{label_exact}")')
    c_tot.font = font_bold
    c_tot.alignment = Alignment(horizontal='center')
    c_pct = ws_resumen.cell(row=idx, column=6, value=f'=E{idx}/E16')
    c_pct.number_format = '0.0%'
    c_pct.alignment = Alignment(horizontal='center')

    fill = ZEBRA_FILL if idx % 2 == 0 else 'FFFFFF'
    for c_i in range(2, 7):
        ws_resumen.cell(row=idx, column=c_i).fill = PatternFill(start_color=fill, end_color=fill, fill_type='solid')
        ws_resumen.cell(row=idx, column=c_i).border = thin_border
    ws_resumen.row_dimensions[idx].height = 22

# Fila Total Tabla 1
ws_resumen.cell(row=16, column=2, value="")
ws_resumen.cell(row=16, column=3, value="TOTAL CATÁLOGO").font = font_bold
ws_resumen.cell(row=16, column=4, value="").font = font_bold
c_tot1 = ws_resumen.cell(row=16, column=5, value=f'=SUM(E11:E15)')
c_tot1.font = font_bold
c_tot1.alignment = Alignment(horizontal='center')
c_pct1 = ws_resumen.cell(row=16, column=6, value=f'=SUM(F11:F15)')
c_pct1.font = font_bold
c_pct1.alignment = Alignment(horizontal='center')
c_pct1.number_format = '0.0%'
for c_i in range(2, 7):
    ws_resumen.cell(row=16, column=c_i).fill = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
    ws_resumen.cell(row=16, column=c_i).border = thin_border

# Tabla 2: Distribución por Frecuencia de Aparición (Polijerarquía)
ws_resumen['B18'] = "2. MODELO DE POLIJERARQUÍA (FRECUENCIA DE APARICIÓN EN EL ÁRBOL)"
ws_resumen['B18'].font = Font(name='Segoe UI', size=11, bold=True, color='1E293B')

t2_headers = ["Frecuencia de Aparición", "Condición Arquitectónica", "Cantidad de Contenidos", "Nodos que Aporta", "% de Contenidos"]
for col_idx, h in enumerate(t2_headers, start=2):
    c = ws_resumen.cell(row=19, column=col_idx, value=h)
    c.font = font_header
    c.fill = PatternFill(start_color=BLUE_HEADER, end_color=BLUE_HEADER, fill_type='solid')
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = header_border
ws_resumen.row_dimensions[19].height = 24

frecuencias_data = [
    ("Aparece 1 sola vez", "Contenido exclusivo de una categoría o rol específico", 1),
    ("Aparece 2 veces", "Contenido compartido entre 2 subramas o segmentos", 2),
    ("Aparece 3 veces", "Herramienta o trámite común para 3 ramas", 3),
    ("Aparece 4 veces", "Recurso transversal para 4 actores de comercio exterior", 4),
    ("Aparece 5 veces", "Recurso transversal para los 5 Auxiliares de la Función Pública", 5),
]

for idx, (frec_label, cond, n_veces) in enumerate(frecuencias_data, start=20):
    ws_resumen.cell(row=idx, column=2, value=frec_label).font = font_bold
    ws_resumen.cell(row=idx, column=3, value=cond)
    c_cant = ws_resumen.cell(row=idx, column=4, value=f'=COUNTIF(\'{SHEET_DATA_NAME}\'!O2:O{LAST_ROW}, {n_veces})')
    c_cant.font = font_bold
    c_cant.alignment = Alignment(horizontal='center')
    c_nodos = ws_resumen.cell(row=idx, column=5, value=f'=D{idx}*{n_veces}')
    c_nodos.alignment = Alignment(horizontal='center')
    c_pct = ws_resumen.cell(row=idx, column=6, value=f'=D{idx}/D25')
    c_pct.number_format = '0.0%'
    c_pct.alignment = Alignment(horizontal='center')

    fill = ZEBRA_FILL if idx % 2 == 0 else 'FFFFFF'
    for c_i in range(2, 7):
        ws_resumen.cell(row=idx, column=c_i).fill = PatternFill(start_color=fill, end_color=fill, fill_type='solid')
        ws_resumen.cell(row=idx, column=c_i).border = thin_border
    ws_resumen.row_dimensions[idx].height = 22

# Fila Total Tabla 2
ws_resumen.cell(row=25, column=2, value="TOTAL CATÁLOGO").font = font_bold
ws_resumen.cell(row=25, column=3, value="Suma de contenidos únicos y total de nodos del árbol").font = font_bold
c_tot2 = ws_resumen.cell(row=25, column=4, value=f'=SUM(D20:D24)')
c_tot2.font = font_bold
c_tot2.alignment = Alignment(horizontal='center')
c_tot_nodos = ws_resumen.cell(row=25, column=5, value=f'=SUM(E20:E24)')
c_tot_nodos.font = font_bold
c_tot_nodos.alignment = Alignment(horizontal='center')
c_pct2 = ws_resumen.cell(row=25, column=6, value=f'=SUM(F20:F24)')
c_pct2.font = font_bold
c_pct2.alignment = Alignment(horizontal='center')
c_pct2.number_format = '0.0%'
for c_i in range(2, 7):
    ws_resumen.cell(row=25, column=c_i).fill = PatternFill(start_color='E2E8F0', end_color='E2E8F0', fill_type='solid')
    ws_resumen.cell(row=25, column=c_i).border = thin_border

resumen_col_widths = [4, 28, 48, 42, 22, 22, 16, 16]
for i, w in enumerate(resumen_col_widths, start=1):
    ws_resumen.column_dimensions[get_column_letter(i)].width = w

# ==============================================================================
# HOJA 2: CATÁLOGO DE CONTENIDOS ÚNICOS (683 REGISTROS)
# ==============================================================================
ws_catalogo = wb.create_sheet(title=SHEET_DATA_NAME)
ws_catalogo.views.sheetView[0].showGridLines = True

headers_catalogo = [
    "No.",
    "Código Canónico Oficial (SAT-GES-####)",
    "ID Content Hub",
    "Nombre de la Gestión / Contenido (Lenguaje Claro)",
    "Nombre Oficial Administrativo",
    "¿Para qué sirve? (Descripción Operativa)",
    "Etapa del Ciclo de Vida",
    "Tipo de Interacción",
    "Tipología de Contenido",
    "Plataforma / Sistema Técnico",
    "Canal de Atención",
    "Perfil Destinatario",
    "Base Legal / Fundamento Jurídico",
    "¿Es Transversal?",
    "Total Audiencias / Nodos",
    "Segmentos Aplicables",
    "¿Es Brecha Normativa?",
    "URL Portal SAT"
]

for col_idx, h in enumerate(headers_catalogo, start=1):
    c = ws_catalogo.cell(row=1, column=col_idx, value=h)
    c.font = font_header
    c.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type='solid')
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = header_border
ws_catalogo.row_dimensions[1].height = 32

for r_idx, item in enumerate(catalogo, start=2):
    es_trans = "SÍ" if item.get('esTransversal') else "NO"
    es_brecha = "SÍ" if item.get('esBrecha') else "NO"
    segs = ", ".join(item.get('segmentosAplicables', []))
    
    ws_catalogo.cell(row=r_idx, column=1, value=r_idx - 1).alignment = Alignment(horizontal='center')
    
    c_cod = ws_catalogo.cell(row=r_idx, column=2, value=item.get('codigo', ''))
    c_cod.font = font_bold
    c_cod.alignment = Alignment(horizontal='center')
    
    c_id = ws_catalogo.cell(row=r_idx, column=3, value=item.get('id', ''))
    c_id.alignment = Alignment(horizontal='center')
    
    c_tit = ws_catalogo.cell(row=r_idx, column=4, value=item.get('titulo', ''))
    c_tit.font = font_bold
    
    ws_catalogo.cell(row=r_idx, column=5, value=item.get('nombreActual', ''))
    ws_catalogo.cell(row=r_idx, column=6, value=item.get('descripcion', ''))
    ws_catalogo.cell(row=r_idx, column=7, value=item.get('etapaAtoLabel', item.get('etapaAto', '')))
    ws_catalogo.cell(row=r_idx, column=8, value=item.get('tipoInteraccionLabel', item.get('tipoInteraccion', '')))
    ws_catalogo.cell(row=r_idx, column=9, value=item.get('tipologiaContenidoLabel', item.get('tipologiaContenido', '')))
    ws_catalogo.cell(row=r_idx, column=10, value=item.get('plataformaSistemaLabel', item.get('plataformaSistema', '')))
    ws_catalogo.cell(row=r_idx, column=11, value=item.get('canalAtencion', ''))
    ws_catalogo.cell(row=r_idx, column=12, value=item.get('perfilDestinatario', ''))
    ws_catalogo.cell(row=r_idx, column=13, value=item.get('baseLegal', ''))
    
    c_trans = ws_catalogo.cell(row=r_idx, column=14, value=es_trans)
    c_trans.alignment = Alignment(horizontal='center')
    if es_trans == "SÍ":
        c_trans.font = font_bold
    
    c_aud = ws_catalogo.cell(row=r_idx, column=15, value=item.get('totalAudiencias', 1))
    c_aud.alignment = Alignment(horizontal='center')
    c_aud.font = font_bold
    
    ws_catalogo.cell(row=r_idx, column=16, value=segs)
    
    c_br = ws_catalogo.cell(row=r_idx, column=17, value=es_brecha)
    c_br.alignment = Alignment(horizontal='center')
    if es_brecha == "SÍ":
        c_br.font = font_bold
    
    # URL con hipervínculo
    url_val = item.get('url', '')
    c_url = ws_catalogo.cell(row=r_idx, column=18, value=url_val)
    if url_val.startswith('http'):
        c_url.font = font_link
        c_url.hyperlink = url_val

    # Color de fila alternado (Zebra) o highlight para transversales
    if item.get('esTransversal'):
        fill_color = HIGHLIGHT_FILL
    elif r_idx % 2 == 0:
        fill_color = ZEBRA_FILL
    else:
        fill_color = 'FFFFFF'

    for c_i in range(1, 19):
        cell = ws_catalogo.cell(row=r_idx, column=c_i)
        if cell.font == font_body or cell.font is None:
            cell.font = font_body
        cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')
        cell.border = thin_border
    
    ws_catalogo.row_dimensions[r_idx].height = 24

catalogo_widths = [6, 22, 22, 40, 36, 50, 24, 24, 24, 24, 16, 28, 42, 16, 16, 26, 16, 50]
for idx, w in enumerate(catalogo_widths, start=1):
    ws_catalogo.column_dimensions[get_column_letter(idx)].width = w

ws_catalogo.freeze_panes = 'C2'
ws_catalogo.auto_filter.ref = f"A1:R{LAST_ROW}"

# ==============================================================================
# HOJA 3: CONTENIDOS TRANSVERSALES (17 REGISTROS POLIJERÁRQUICOS)
# ==============================================================================
ws_trans = wb.create_sheet(title="Contenidos Transversales (17)")
ws_trans.views.sheetView[0].showGridLines = True

trans_headers = [
    "No.",
    "Código Oficial",
    "Nombre del Contenido Canónico",
    "Apariciones",
    "Segmentos que lo Comparten",
    "Ubicaciones / Nodos en el Árbol de Navegación",
    "¿Para qué sirve?",
    "URL Portal SAT"
]

for col_idx, h in enumerate(trans_headers, start=1):
    c = ws_trans.cell(row=1, column=col_idx, value=h)
    c.font = font_header
    c.fill = PatternFill(start_color=PURPLE_ACCENT, end_color=PURPLE_ACCENT, fill_type='solid')
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = header_border
ws_trans.row_dimensions[1].height = 30

transversales_items = [c for c in catalogo if c.get('esTransversal')]

for r_idx, item in enumerate(transversales_items, start=2):
    # Formatear lista de ubicaciones
    ubics = item.get('ubicaciones', [])
    ubics_text = "\n".join([f"• {u.get('segmento', '')} > {u.get('categoria', '')} > {u.get('subcategoria', '')}" for u in ubics])
    segs = ", ".join(item.get('segmentosAplicables', []))

    ws_trans.cell(row=r_idx, column=1, value=r_idx - 1).alignment = Alignment(horizontal='center')
    
    c_cod = ws_trans.cell(row=r_idx, column=2, value=item.get('codigo', ''))
    c_cod.font = font_bold
    c_cod.alignment = Alignment(horizontal='center')
    
    c_tit = ws_trans.cell(row=r_idx, column=3, value=item.get('titulo', ''))
    c_tit.font = font_bold
    
    c_aud = ws_trans.cell(row=r_idx, column=4, value=item.get('totalAudiencias', 1))
    c_aud.font = font_bold
    c_aud.alignment = Alignment(horizontal='center')
    
    ws_trans.cell(row=r_idx, column=5, value=segs).font = font_bold
    
    c_ub = ws_trans.cell(row=r_idx, column=6, value=ubics_text)
    c_ub.alignment = Alignment(wrap_text=True)
    
    ws_trans.cell(row=r_idx, column=7, value=item.get('descripcion', ''))
    
    url_val = item.get('url', '')
    c_url = ws_trans.cell(row=r_idx, column=8, value=url_val)
    if url_val.startswith('http'):
        c_url.font = font_link
        c_url.hyperlink = url_val

    fill_color = 'F5F3FF' if r_idx % 2 == 0 else 'FFFFFF'
    for c_i in range(1, 9):
        cell = ws_trans.cell(row=r_idx, column=c_i)
        if cell.font is None: cell.font = font_body
        cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')
        cell.border = thin_border
    
    # Altura dinámica según cantidad de ubicaciones
    ws_trans.row_dimensions[r_idx].height = max(26, len(ubics) * 16)

trans_widths = [6, 18, 38, 14, 26, 52, 45, 45]
for idx, w in enumerate(trans_widths, start=1):
    ws_trans.column_dimensions[get_column_letter(idx)].width = w

ws_trans.freeze_panes = 'C2'
ws_trans.auto_filter.ref = f"A1:H{len(transversales_items)+1}"

# ==============================================================================
# HOJA 4: MATRIZ DE MAPEO A NODOS DE AUDIENCIA (715 APARICIONES)
# ==============================================================================
ws_nodos = wb.create_sheet(title="Mapeo a Nodos (715)")
ws_nodos.views.sheetView[0].showGridLines = True

nodos_headers = [
    "No.",
    "Código Oficial (SAT-GES-####)",
    "ID Nodo de Navegación",
    "Nombre de la Gestión / Contenido",
    "Segmento (Nivel 1)",
    "Categoría (Nivel 2)",
    "Subcategoría (Nivel 3)",
    "Tema / Subárea (Nivel 4)",
    "Ruta de Navegación (Miga de Pan)",
    "Etapa del Ciclo de Vida",
    "Tipo de Interacción",
    "URL Portal SAT"
]

for col_idx, h in enumerate(nodos_headers, start=1):
    c = ws_nodos.cell(row=1, column=col_idx, value=h)
    c.font = font_header
    c.fill = PatternFill(start_color=CYAN_ACCENT, end_color=CYAN_ACCENT, fill_type='solid')
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = header_border
ws_nodos.row_dimensions[1].height = 30

row_nodo = 2
for item in catalogo:
    ubics = item.get('ubicaciones', [])
    for u in ubics:
        ws_nodos.cell(row=row_nodo, column=1, value=row_nodo - 1).alignment = Alignment(horizontal='center')
        
        c_cod = ws_nodos.cell(row=row_nodo, column=2, value=item.get('codigo', ''))
        c_cod.font = font_bold
        c_cod.alignment = Alignment(horizontal='center')
        
        c_tid = ws_nodos.cell(row=row_nodo, column=3, value=u.get('tramiteId', ''))
        c_tid.alignment = Alignment(horizontal='center')
        
        c_tit = ws_nodos.cell(row=row_nodo, column=4, value=item.get('titulo', ''))
        c_tit.font = font_bold
        
        ws_nodos.cell(row=row_nodo, column=5, value=u.get('segmento', '')).font = font_bold
        ws_nodos.cell(row=row_nodo, column=6, value=u.get('categoria', ''))
        ws_nodos.cell(row=row_nodo, column=7, value=u.get('subcategoria', ''))
        ws_nodos.cell(row=row_nodo, column=8, value=u.get('tema', u.get('subtema', '—')))
        ws_nodos.cell(row=row_nodo, column=9, value=u.get('migaBreadcrumb', ''))
        ws_nodos.cell(row=row_nodo, column=10, value=item.get('etapaAtoLabel', item.get('etapaAto', '')))
        ws_nodos.cell(row=row_nodo, column=11, value=item.get('tipoInteraccionLabel', item.get('tipoInteraccion', '')))
        
        url_val = item.get('url', '')
        c_url = ws_nodos.cell(row=row_nodo, column=12, value=url_val)
        if url_val.startswith('http'):
            c_url.font = font_link
            c_url.hyperlink = url_val

        fill_color = ZEBRA_FILL if row_nodo % 2 == 0 else 'FFFFFF'
        for c_i in range(1, 13):
            cell = ws_nodos.cell(row=row_nodo, column=c_i)
            if cell.font is None: cell.font = font_body
            cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type='solid')
            cell.border = thin_border
        
        ws_nodos.row_dimensions[row_nodo].height = 22
        row_nodo += 1

TOTAL_NODOS_ROW = row_nodo - 1
nodos_widths = [6, 18, 18, 38, 22, 28, 28, 28, 48, 22, 22, 45]
for idx, w in enumerate(nodos_widths, start=1):
    ws_nodos.column_dimensions[get_column_letter(idx)].width = w

ws_nodos.freeze_panes = 'D2'
ws_nodos.auto_filter.ref = f"A1:L{TOTAL_NODOS_ROW}"

# ==============================================================================
# GUARDAR LIBRO EXCEL EN ENTREGABLES
# ==============================================================================
OUT_ENTREGABLE = 'docs/fuentes-datos/02_archivos_producto_entregables/Catalogo_Contenidos_Unicos_Portal_SAT.xlsx'
os.makedirs(os.path.dirname(OUT_ENTREGABLE), exist_ok=True)

try:
    wb.save(OUT_ENTREGABLE)
    print(f"Libro guardado exitosamente en: {OUT_ENTREGABLE}")
except Exception as e:
    print(f"Aviso al guardar en entregables: {e}")

print(f"Hojas creadas: {', '.join(wb.sheetnames)}")
