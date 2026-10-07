import openpyxl

excel_path = 'c:/Users/busqu/Documents/GitHub/SAT/docs/Arbol_de_Navegacion_Portal_v5.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)
print("Hojas en Arbol_de_Navegacion_Portal_v5.xlsx:", wb.sheetnames)

for s in wb.sheetnames:
    ws = wb[s]
    print(f"\nHoja '{s}': {ws.max_row} filas, {ws.max_column} columnas")
    header = [ws.cell(row=1, column=c).value for c in range(1, min(ws.max_column+1, 15))]
    print("  Header:", header[:10])
    # Buscar si hay mención de depósitos o almacenes
    count = 0
    for r in range(2, min(ws.max_row+1, 500)):
        row_str = " ".join([str(ws.cell(row=r, column=c).value) for c in range(1, 10)]).lower()
        if 'depósito' in row_str or 'deposito' in row_str or 'almacén' in row_str or 'almacen' in row_str:
            count += 1
            if count <= 5:
                vals = [ws.cell(row=r, column=c).value for c in range(1, 8)]
                print(f"    Match {count}: {vals}")
    print(f"  Total matches en {s}: {count}")
