import json
import openpyxl
import os

print("=== INICIANDO HOMOLOGACIÓN Y SANEAMIENTO DE TIPOLOGÍA DE CONTENIDO ===")

# Mapeo canónico de etiquetas
MAP_TIPOLOGIA_LABEL = {
    'tramite_interactivo': 'Trámite / Aplicativo en Línea',
    'guia_requisitos': 'Guía de Requisitos y Pasos',
    'guia_informativa': 'Guía Informativa / Texto',
    'consulta_buscador': 'Consulta Pública / Buscador',
    'contenido_multimedia': 'Multimedia / Curso Virtual',
    'normativa_criterio': 'Normativa y Criterio Legal',
    'recurso_descargable': 'Recurso / Software Descargable',
    'aviso_operativo': 'Aviso Operativo'
}

MAP_INTERACCION_LABEL = {
    'servicio_transaccional': 'Trámite / Aplicativo en Línea',
    'consulta_datos': 'Consulta en Base de Datos',
    'guia_informativa': 'Guía Informativa / Texto',
    'descarga_recurso': 'Descarga de Documento / Software'
}

# 1. ACTUALIZAR allTramites.json
json_path = 'src/data/allTramites.json'
with open(json_path, 'r', encoding='utf-8') as f:
    tramites = json.load(f)

# IDs específicos con correcciones semánticas directas
ESPECIFICOS = {
    'SAT-GES-0008': { # Verificar títulos con código QR
        'tipologiaContenido': 'consulta_buscador',
        'tipoInteraccion': 'consulta_datos'
    },
    'SAT-GES-0039': { # Ver el estado de tu cita
        'tipologiaContenido': 'consulta_buscador',
        'tipoInteraccion': 'consulta_datos'
    },
    'SAT-GES-0065': { # Imprimir tu constancia del RTU
        'tipologiaContenido': 'consulta_buscador',
        'tipoInteraccion': 'consulta_datos'
    },
    'SAT-GES-0207': { # Ver el estado de tu queja
        'tipologiaContenido': 'consulta_buscador',
        'tipoInteraccion': 'consulta_datos'
    },
    'SAT-GES-0333': { # Historial de facilidades de pago
        'tipologiaContenido': 'consulta_buscador',
        'tipoInteraccion': 'consulta_datos'
    },
    'SAT-GES-0421': { # Instalador y Soporte de Firma Digital ActiveX PKI / DUA
        'tipologiaContenido': 'recurso_descargable',
        'tipoInteraccion': 'descarga_recurso'
    }
}

homologados_transaccional = 0
for t in tramites:
    cid = t.get('codigo_oficial') or t.get('id')

    # A) Aplicar correcciones específicas si corresponde
    if cid in ESPECIFICOS:
        t['tipologiaContenido'] = ESPECIFICOS[cid]['tipologiaContenido']
        t['tipoInteraccion'] = ESPECIFICOS[cid]['tipoInteraccion']

    # B) Homologar tramite_transaccional -> tramite_interactivo
    if t.get('tipologiaContenido') == 'tramite_transaccional':
        t['tipologiaContenido'] = 'tramite_interactivo'
        homologados_transaccional += 1

    # C) Asegurar labels sincronizados
    tc = t.get('tipologiaContenido')
    if tc in MAP_TIPOLOGIA_LABEL:
        t['tipologiaContenidoLabel'] = MAP_TIPOLOGIA_LABEL[tc]

    ti = t.get('tipoInteraccion')
    if ti in MAP_INTERACCION_LABEL:
        t['tipoInteraccionLabel'] = MAP_INTERACCION_LABEL[ti]

with open(json_path, 'w', encoding='utf-8') as f:
    json.dump(tramites, f, ensure_ascii=False, indent=2)

print(f"allTramites.json saneado. {homologados_transaccional} trámites transaccionales homologados.")

# 2. ACTUALIZAR database/sat_portal_nosql.json si existe
nosql_path = 'database/sat_portal_nosql.json'
if os.path.exists(nosql_path):
    with open(nosql_path, 'r', encoding='utf-8') as f:
        nosql = json.load(f)
    if 'tramites' in nosql:
        for t in nosql['tramites']:
            cid = t.get('codigo_oficial') or t.get('id')
            if cid in ESPECIFICOS:
                t['tipologia'] = ESPECIFICOS[cid]['tipologiaContenido']
                t['tipo_interaccion'] = ESPECIFICOS[cid]['tipoInteraccion']
            if t.get('tipologia') == 'tramite_transaccional':
                t['tipologia'] = 'tramite_interactivo'
        with open(nosql_path, 'w', encoding='utf-8') as f:
            json.dump(nosql, f, ensure_ascii=False, indent=2)
        print("sat_portal_nosql.json actualizado con éxito.")

# 3. ACTUALIZAR EXCEL MAESTRO
excel_path = 'docs/fuentes-datos/02_archivos_producto_entregables/Estructura_Final_Contenido_Portal_SAT_Actualizado.xlsx'
wb = openpyxl.load_workbook(excel_path)

# Crear un mapa rápido desde el JSON ya saneado
dict_by_id = { (t.get('codigo_oficial') or t.get('id')): t for t in tramites }

sheets_to_update = ['Matriz Maestra (715)', 'Contribuyentes (344)', 'Comercio Exterior (245)', 'Profesionales (47)', 'Entes Exentos (79)']

for sheet_name in sheets_to_update:
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers = [cell.value for cell in ws[1]]
        if 'ID Gestión' in headers and 'Tipología de Contenido' in headers and 'Tipo de Interacción' in headers:
            col_id = headers.index('ID Gestión') + 1
            col_tc = headers.index('Tipología de Contenido') + 1
            col_ti = headers.index('Tipo de Interacción') + 1

            for row_idx in range(2, ws.max_row + 1):
                val_id = ws.cell(row=row_idx, column=col_id).value
                if val_id in dict_by_id:
                    item = dict_by_id[val_id]
                    ws.cell(row=row_idx, column=col_tc, value=item['tipologiaContenidoLabel'])
                    ws.cell(row=row_idx, column=col_ti, value=item['tipoInteraccionLabel'])

wb.save(excel_path)
print(f"Libro {excel_path} sincronizado y guardado con éxito.")
print("=== SANEAMIENTO DE TIPOLOGÍAS FINALIZADO ===")
