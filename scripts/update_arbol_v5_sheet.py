import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json
from assign_base_legal import get_base_legal

wb = openpyxl.load_workbook('docs/fuentes-datos/Arbol_de_Navegacion_Portal_v5.xlsx')
sheet = wb['Operadores de comercio exterior']

# Cargar los 202 registros afinados directamente de allTramites.json
all_tramites = json.load(open('src/data/allTramites.json', encoding='utf-8'))
recs = [t for t in all_tramites if t.get('pillar') == 'comercio_exterior']

# Limpiar filas existentes en la hoja desde la fila 4
for r in range(sheet.max_row, 3, -1):
    sheet.delete_rows(r)

# Encabezados en Fila 4 con Base Legal explícita
headers = [
    'Macro-Grupo',
    'Actor / Rol',
    'Especialidad / Tema',
    'Área Funcional (Subtema)',
    'Nombre del Trámite',
    'Descripción (Lenguaje Ciudadano Activo)',
    'Base Legal (En base a qué: CAUCA, RECAUCA, Leyes)',
    'Control de Auditoría',
    'Etapa Ciclo ATO',
    'Tipo de Interacción',
    'Es Brecha',
    'Ruta de Navegación (Miga de Pan)',
    'URL Oficial SAT',
    'Notas / Justificación'
]

# Estilos SAT
header_fill = PatternFill(start_color='14649B', end_color='14649B', fill_type='solid')
header_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
data_font = Font(name='Segoe UI', size=9)
bold_font = Font(name='Segoe UI', size=9, bold=True)
url_font = Font(name='Segoe UI', size=9, color='0284C7', underline='single')
zebra_fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
brecha_fill = PatternFill(start_color='FFF7ED', end_color='FFF7ED', fill_type='solid')
brecha_font = Font(name='Segoe UI', size=9, bold=True, color='C25E00')

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

# Escribir encabezados en fila 4
for col_idx, h in enumerate(headers, 1):
    c = sheet.cell(row=4, column=col_idx, value=h)
    c.fill = header_fill
    c.font = header_font
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = thin_border

sheet.row_dimensions[4].height = 28

# Escribir los 202 registros a partir de la fila 5
for row_idx, r in enumerate(recs, 5):
    base_leg = get_base_legal(r)
    
    sheet.cell(row=row_idx, column=1, value=r['macroGrupo'])
    sheet.cell(row=row_idx, column=2, value=r['categoria'])
    sheet.cell(row=row_idx, column=3, value=r['tema'])
    sheet.cell(row=row_idx, column=4, value=r['subcategoria'])
    sheet.cell(row=row_idx, column=5, value=r['tramite'])
    sheet.cell(row=row_idx, column=6, value=r['descripcion'])
    sheet.cell(row=row_idx, column=7, value=base_leg)
    
    # Control Auditoría
    c_audit = sheet.cell(row=row_idx, column=8, value=r['controlAuditoria'])
    if r['esBrecha']:
        c_audit.font = brecha_font
        c_audit.fill = brecha_fill
    else:
        c_audit.font = bold_font
    
    sheet.cell(row=row_idx, column=9, value=r['etapaAtoLabel'])
    sheet.cell(row=row_idx, column=10, value=r['tipoInteraccionLabel'])
    sheet.cell(row=row_idx, column=11, value='SÍ' if r['esBrecha'] else 'NO')
    sheet.cell(row=row_idx, column=12, value=r['migaBreadcrumb'])
    
    # URL Oficial
    c_url = sheet.cell(row=row_idx, column=13, value=r['url'])
    if r['url']:
        c_url.hyperlink = r['url']
        c_url.font = url_font
    
    sheet.cell(row=row_idx, column=14, value=r['nota'])

    is_even = (row_idx % 2 == 0)
    for c_idx in range(1, 15):
        cell = sheet.cell(row=row_idx, column=c_idx)
        cell.border = thin_border
        if cell.column != 8: # preserve audit formatting
            cell.font = url_font if c_idx == 13 else (bold_font if c_idx == 7 else data_font)
            if is_even and not (r['esBrecha'] and c_idx == 8):
                cell.fill = zebra_fill
        
        # Alignment
        if c_idx in [8, 9, 10, 11]:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        elif c_idx == 13:
            cell.alignment = Alignment(horizontal='left', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=(c_idx in [5, 6, 7]))

    sheet.row_dimensions[row_idx].height = 24

# Auto-ajustar anchos
col_widths = {
    1: 28,  # Macro-Grupo
    2: 32,  # Actor / Rol
    3: 24,  # Tema
    4: 24,  # Subtema
    5: 42,  # Trámite
    6: 55,  # Descripción
    7: 48,  # Base Legal (CAUCA, RECAUCA, Leyes)
    8: 24,  # Control Auditoría
    9: 24,  # Etapa ATO
    10: 26, # Tipo Interacción
    11: 12, # Es Brecha
    12: 45, # Breadcrumb
    13: 45, # URL
    14: 35  # Notas
}

for col_idx, width in col_widths.items():
    col_letter = get_column_letter(col_idx)
    sheet.column_dimensions[col_letter].width = width

sheet.freeze_panes = 'E5'

wb.save('docs/fuentes-datos/Arbol_de_Navegacion_Portal_v5.xlsx')
print("Éxito: Hoja 'Operadores de comercio exterior' en Arbol_de_Navegacion_Portal_v5.xlsx actualizada con Base Legal (CAUCA/RECAUCA).")
